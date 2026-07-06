import openpyxl

from bank_statement_analyzer.calculators.calculator import (
    all_account_continuity,
    expense_by_category,
    group_income_by_counterparty,
    run_totals,
)
from bank_statement_analyzer.classification.classifier import (
    detect_anomalies,
    detect_recurring,
    tag_transactions,
)
from bank_statement_analyzer.export.csv_export import export_csv
from bank_statement_analyzer.export.xlsx_export import export_workbook
from bank_statement_analyzer.parsing.models import StatementMeta
from bank_statement_analyzer.reconciliation.refund_matcher import match_refunds
from bank_statement_analyzer.reconciliation.transfer_matcher import match_self_transfers
from bank_statement_analyzer.run_context import NeedsAttentionItem, ReconciliationResult, RunResult


def _build_run(make_txn, tmp_path):
    txns = [
        # self-transfer pair across two accounts
        make_txn(1, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-05", "UPI/DR/rrn111111111/self xfer", debit=10000),
        make_txn(1, "SBI", "sbi.pdf", "ACC-2", "2025-04-05", "UPI/CR/rrn111111111/self xfer", credit=10000),
        # refund pair, same account
        make_txn(2, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-10", "Flipkart order payment", debit=2500),
        make_txn(3, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-15", "Flipkart refund reversed", credit=2500),
        # salary income
        make_txn(4, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-01", "ACME CORP SALARY CREDIT", credit=80000),
        make_txn(5, "HDFC", "hdfc.pdf", "ACC-1", "2025-05-01", "ACME CORP SALARY CREDIT", credit=80000),
        # CC payment
        make_txn(6, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-20", "CC PAYMENT AUTOPAY", debit=15000),
        # cash withdrawal expense
        make_txn(7, "HDFC", "hdfc.pdf", "ACC-1", "2025-04-21", "ATM CASH WDL", debit=3000),
    ]

    self_transfers = match_self_transfers(txns)
    refunds = match_refunds(txns)
    excluded = set()
    for pair in (*self_transfers.confirmed, *refunds.confirmed):
        excluded.add(pair.leg_a_ref)
        excluded.add(pair.leg_b_ref)

    tag_transactions(txns, excluded_refs=excluded)
    for t in txns:
        if t.ref not in excluded and "SALARY" in t.description:
            t.itr_head = "Business/Professional receipts"
            t.itr_head_confidence = "medium"
            t.counterparty = "ACME CORP"

    detect_recurring(txns, excluded_refs=excluded)
    detect_anomalies(txns, excluded_refs=excluded)

    credits = [t for t in txns if t.is_credit and t.ref not in excluded]
    debits = [t for t in txns if t.is_debit and t.ref not in excluded]
    income_groups = group_income_by_counterparty(credits)
    expense_categories = expense_by_category(debits)
    cc_refs = [t.ref for t in debits if "cc_payment_or_fee" in t.tags]

    metas = {
        "ACC-1": StatementMeta(bank="HDFC", pdf="hdfc.pdf", account_ref="ACC-1", opening_balance=100000, closing_balance=100000 + 80000 + 80000 + 2500 - 10000 - 2500 - 15000 - 3000),
        "ACC-2": StatementMeta(bank="SBI", pdf="sbi.pdf", account_ref="ACC-2", opening_balance=5000, closing_balance=15000),
    }
    continuity = all_account_continuity(txns, metas)

    run = RunResult(
        client_label="Test Client",
        statement_metas=metas,
        transactions=txns,
        self_transfers=ReconciliationResult(confirmed=self_transfers.confirmed, unmatched=self_transfers.unmatched),
        refunds=ReconciliationResult(confirmed=refunds.confirmed, unmatched=refunds.unmatched),
        income_groups=income_groups,
        expense_categories=expense_categories,
        cc_payment_refs=cc_refs,
        run_totals=run_totals(txns),
        continuity=continuity,
        needs_attention=[NeedsAttentionItem(issue="test flag", file="hdfc.pdf", ref=txns[0].ref, suggestion="review")],
    )
    return run


def test_workbook_has_all_sheets_and_refunds_sheet(make_txn, tmp_path):
    run = _build_run(make_txn, tmp_path)
    out_path = tmp_path / "output.xlsx"
    export_workbook(run, str(out_path))

    wb = openpyxl.load_workbook(str(out_path))
    expected_sheets = {
        "Summary", "Self-transfers", "Refunds", "Income sources",
        "Credit-card payments", "Expenditure summary", "Needs attention",
        "All transactions",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))

    refunds_ws = wb["Refunds"]
    header = [c.value for c in next(refunds_ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["Original debit ref", "Refund credit ref", "Amount", "Date diff (days)", "Method", "Confidence", "Status"]
    data_row = [c.value for c in next(refunds_ws.iter_rows(min_row=2, max_row=2))]
    assert data_row[2] == 2500
    assert data_row[6] == "confirmed"

    self_ws = wb["Self-transfers"]
    self_data_row = [c.value for c in next(self_ws.iter_rows(min_row=2, max_row=2))]
    assert self_data_row[2] == 10000


def test_balance_continuity_ok_in_workbook(make_txn, tmp_path):
    run = _build_run(make_txn, tmp_path)
    for c in run.continuity:
        assert c.ok, f"{c.account_ref} mismatch: {c.diff}"


def test_csv_export_mirrors_all_transactions(make_txn, tmp_path):
    run = _build_run(make_txn, tmp_path)
    out_path = tmp_path / "output.csv"
    export_csv(run, str(out_path))
    content = out_path.read_text(encoding="utf-8-sig")
    lines = content.strip().splitlines()
    assert lines[0].startswith("Serial,Bank,PDF,Date,Description")
    assert len(lines) == len(run.transactions) + 1
